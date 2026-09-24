# -*- coding: utf-8 -*-
"""La hoja para Recursos Humanos: el personal de seguridad y lo que le falta.

Pedido de Salvador, 23 de septiembre. En vez de que RH abra ficha por ficha
en Odoo, esta hoja trae a las personas que van a entrar a la operacion
--puesto escrito «Personal de Seguridad...» o «Security Driver...»-- con
lo que ya tienen en Odoo y lo que les falta marcado en amarillo. RH llena
las casillas y la hoja se carga de regreso a Odoo con otro script, con
ensayo y bitacora para deshacer, igual que ordenar_odoo.py.

SOLO LEE Odoo. La hoja lleva datos personales: se escribe en la carpeta
«Claude outputs», que no va a git, y en la terminal solo salen conteos.

Desde la raiz del proyecto:

    docker compose run --rm -v "$HOME/Desktop/centauro/Claude outputs:/salida" \\
      api sh -c "pip install -q openpyxl && python hoja_rh_odoo.py"
"""
import collections
import os
import sys
import unicodedata
from datetime import datetime

import httpx

PERMITIDOS = {"search_read", "fields_get"}
PUESTOS = ("personal de seguridad", "security driver")
PLAZAS = ["Ciudad de México", "Guadalajara", "Querétaro", "Monterrey"]
# Los errores de dedo que ya aparecieron o que aparecen siempre.
DOMINIOS_RAROS = {"gamil.com", "gmial.com", "gmai.com", "gmail.con", "gmail.co",
                  "hotmial.com", "hotmal.com", "hotmail.con", "hotamil.com",
                  "yaho.com", "yahoo.con", "outlok.com", "outlook.con"}
# Como empieza en base64 una foto de verdad: PNG, JPEG, GIF o WEBP. El
# circulo con iniciales que Odoo le pone a quien no tiene foto es un SVG,
# y no cuenta.
FOTOS = ("iVBOR", "/9j/", "R0lGOD", "UklGR")


class Odoo:
    def __init__(self, base, llave):
        base = base.strip().rstrip("/")
        self.base = base if base.startswith("http") else "https://" + base
        self.http = httpx.Client(timeout=60, headers={
            "Authorization": f"bearer {llave}",
            "Content-Type": "application/json; charset=utf-8",
            "User-Agent": "centauro-hoja-rh/1.0"})

    def llamar(self, modelo, metodo, **args):
        if metodo not in PERMITIDOS:
            raise RuntimeError(f"{metodo} no es de solo lectura")
        r = self.http.post(f"{self.base}/json/2/{modelo}/{metodo}",
                           json={"context": {"lang": "es_MX"}, **args})
        if r.status_code != 200:
            try:
                mensaje = r.json().get("message") or r.text
            except ValueError:
                mensaje = r.text
            raise RuntimeError(f"{r.status_code}: {' '.join(str(mensaje).split())[:200]}")
        return r.json()


def normal(texto):
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().split())


def nombre_de(valor):
    if isinstance(valor, (list, tuple)) and len(valor) > 1:
        return str(valor[1])
    if isinstance(valor, dict):
        return str(valor.get("display_name") or valor.get("name") or "")
    return ""


def texto(valor):
    return "" if valor in (False, None) else str(valor).strip()


def main():
    base = os.environ.get("ODOO_BASE", "").strip()
    llave = os.environ.get("ODOO_API_KEY", "").strip()
    if not base or not llave:
        print("Falta ODOO_BASE u ODOO_API_KEY en el .env.")
        return 1
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("Falta openpyxl: corre el comando con «pip install -q openpyxl &&».")
        return 1

    odoo = Odoo(base, llave)
    campos = odoo.llamar("hr.employee", "fields_get", attributes=["type"])
    opcionales = ["work_location_id", "mobile_phone", "work_email",
                  "private_email", "registration_number", "image_128"]
    leer = ["name", "job_title"] + [c for c in opcionales if c in campos]
    try:
        empleados = odoo.llamar("hr.employee", "search_read", domain=[],
                                fields=leer, order="name")
    except RuntimeError as error:
        print(f"No se pudo leer Odoo: {error}")
        return 3
    gente = [e for e in empleados
             if normal(e.get("job_title")).startswith(PUESTOS)]
    if not gente:
        print("No se encontro a nadie con puesto «Personal de Seguridad» o "
              "«Security Driver». No se hizo la hoja: revisa en Odoo.")
        return 4

    # ------------------------------------------------ los datos, fila por fila
    correos = collections.Counter()
    for e in gente:
        for c in ("work_email", "private_email"):
            if texto(e.get(c)):
                correos[texto(e.get(c)).lower()] += 1

    filas, falta_total, revisar = [], collections.Counter(), 0
    for e in gente:
        plaza = nombre_de(e.get("work_location_id"))
        plaza = next((p for p in PLAZAS if normal(p) == normal(plaza)), "")
        trabajo, personal = texto(e.get("work_email")), texto(e.get("private_email"))
        foto = texto(e.get("image_128")).startswith(FOTOS)
        falta = []
        if not plaza:
            falta.append("plaza")
        if not texto(e.get("mobile_phone")):
            falta.append("celular")
        if not texto(e.get("registration_number")):
            falta.append("referencia")
        # Con el personal entra a la app: los de trabajo del personal de
        # seguridad se van a suspender.
        if not personal:
            falta.append("correo personal")
        if not foto:
            falta.append("foto")
        dudosos = [c for c in (trabajo, personal) if c and (
            c.split("@")[-1].lower() in DOMINIOS_RAROS
            or correos[c.lower()] > 1 or "@" not in c or " " in c)]
        if dudosos:
            falta.append("revisar correo")
            revisar += 1
        falta_total.update(falta)
        filas.append({"id": e["id"], "nombre": texto(e.get("name")),
                      "puesto": texto(e.get("job_title")), "plaza": plaza,
                      "celular": texto(e.get("mobile_phone")),
                      "trabajo": trabajo, "personal": personal,
                      "referencia": texto(e.get("registration_number")),
                      "foto": "Sí" if foto else "No",
                      "falta": ", ".join(falta), "dudosos": dudosos})

    # ------------------------------------------------ la hoja
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Personal"
    NAVY = "1B1546"
    cab = Font(bold=True, color="FFFFFF")
    fondo_cab = PatternFill("solid", fgColor=NAVY)
    gris = PatternFill("solid", fgColor="EEF1F4")
    amarillo = PatternFill("solid", fgColor="FFF2A8")
    naranja = PatternFill("solid", fgColor="FFD5A8")

    columnas = [
        ("No. Odoo", "id", 10, False),
        ("Nombre", "nombre", 32, False),
        ("Puesto escrito", "puesto", 30, False),
        ("Plaza", "plaza", 18, True),
        ("Celular de trabajo", "celular", 18, True),
        ("Correo de trabajo", "trabajo", 30, True),
        ("Correo personal", "personal", 30, True),
        ("Referencia de empleado", "referencia", 16, True),
        ("¿Tiene foto en Odoo?", "foto", 12, False),
        ("Falta", "falta", 30, False),
        ("Notas de RH", "notas", 34, True),
    ]
    for i, (titulo, _, ancho, _) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=i, value=titulo)
        celda.font, celda.fill = cab, fondo_cab
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.row_dimensions[1].height = 32
    hoja["A1"].comment = Comment("No cambiar: es con lo que se encuentra a "
                                 "cada persona en Odoo.", "Centauro")

    for r, fila in enumerate(filas, start=2):
        for c, (_, clave, _, editable) in enumerate(columnas, start=1):
            valor = fila.get(clave, "")
            celda = hoja.cell(row=r, column=c, value=valor if valor != "" else None)
            if clave in ("celular", "referencia"):
                celda.number_format = "@"          # que Excel no lo vuelva numero
            if not editable:
                celda.fill = gris
            elif clave in ("plaza", "celular", "referencia") and not valor:
                celda.fill = amarillo
            elif clave in ("trabajo", "personal") and valor in fila["dudosos"]:
                celda.fill = naranja
            elif clave == "personal" and not fila["personal"]:
                celda.fill = amarillo
        if fila["foto"] == "No":
            hoja.cell(row=r, column=9).fill = amarillo

    ultima = len(filas) + 1
    lista = DataValidation(type="list", formula1='"' + ",".join(PLAZAS) + '"',
                           allow_blank=True, showErrorMessage=True,
                           errorTitle="Plaza", error="Escoge una de las cuatro plazas.")
    hoja.add_data_validation(lista)
    lista.add(f"D2:D{max(ultima, 2)}")
    hoja.freeze_panes = "D2"
    hoja.auto_filter.ref = f"A1:K{max(ultima, 1)}"

    guia = libro.create_sheet("Instrucciones")
    renglones = [
        ("Personal de seguridad · datos para Odoo", True),
        (f"Hecha el {datetime.now():%d/%m/%Y} desde Odoo. {len(filas)} personas.", False),
        ("", False),
        ("Qué hacer", True),
        ("1. Llenar las casillas en amarillo: plaza, celular de trabajo, correo personal y referencia de empleado.", False),
        ("2. Revisar las casillas en naranja: son correos con un error de dedo probable o repetidos.", False),
        ("3. Cada persona necesita su correo personal: con él entra a la app de campo. "
         "Los correos de trabajo del personal de seguridad se van a suspender.", False),
        ("4. Si algo ya estaba pero está mal, se corrige aquí mismo, encima.", False),
        ("5. La foto no va en esta hoja: se sube en Odoo, en la ficha de la persona. "
         "El círculo con iniciales no es foto.", False),
        ("6. Guardar la hoja con el mismo nombre y avisar a Dirección: se carga a Odoo de un jalón.", False),
        ("", False),
        ("No cambiar", True),
        ("Las columnas en gris: No. Odoo, Nombre, Puesto escrito, Foto y Falta. El No. Odoo es con lo que se encuentra a cada persona.", False),
        ("", False),
        ("Plazas", True),
        ("Ciudad de México, Guadalajara, Querétaro o Monterrey. El Estado de México va como Ciudad de México.", False),
    ]
    for r, (valor, negrita) in enumerate(renglones, start=1):
        celda = guia.cell(row=r, column=1, value=valor)
        celda.font = Font(bold=negrita, size=13 if r == 1 else 11,
                          color=NAVY if negrita else "000000")
    guia.column_dimensions["A"].width = 120
    libro.move_sheet("Instrucciones", offset=-1)
    libro.active = 1

    carpeta = "/salida" if os.path.isdir("/salida") else os.path.dirname(os.path.abspath(__file__))
    nombre = f"Personal_de_seguridad_para_RH_{datetime.now():%Y%m%d}.xlsx"
    libro.save(os.path.join(carpeta, nombre))

    print(f"Hoja lista: {nombre}, con {len(filas)} personas.")
    print("Faltan: " + " · ".join(f"{k} {falta_total[k]}" for k in
                                   ("plaza", "celular", "referencia", "correo personal", "foto")))
    print(f"Personas con un correo que revisar: {revisar}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
