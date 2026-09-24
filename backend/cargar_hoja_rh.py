# -*- coding: utf-8 -*-
"""Cargar a Odoo la hoja que devuelve Recursos Humanos.

La hoja la arma hoja_rh_odoo.py: el personal de seguridad con lo que le
falta en amarillo. RH llena plaza, celular de trabajo, correos y
referencia de empleado, y la devuelve. Esto la sube a Odoo de un jalon,
en vez de ficha por ficha. Pedido de Salvador, 23 de septiembre.

Escribe SOLO en hr.employee y SOLO cinco campos: la ubicacion de trabajo
(la plaza), el celular de trabajo, los dos correos y la referencia de
empleado. No crea ni borra nada.

  * A cada persona se le encuentra por el No. Odoo, y ademas el nombre de
    la fila tiene que ser el de Odoo: una fila movida no le escribe a
    otra persona. Si no coincide, la fila entera se salta, y la persona
    no se da por vista: si su fila buena viene despues, se toma.
  * Una casilla vacia no borra lo que Odoo ya tiene.
  * Lo que no pasa la revision --una plaza que no es de las cuatro, un
    celular de menos de diez digitos, un correo mal escrito o repetido,
    una referencia repetida-- no se escribe y se cuenta; lo demas de esa
    persona si.
  * Solo el personal de seguridad: la fila de alguien mas se salta.

Tres modos, desde la raiz del proyecto, con la hoja en «Claude outputs»:

    docker compose run --rm -v "$HOME/Desktop/centauro/Claude outputs:/salida" \\
      api sh -c "pip install -q openpyxl && python cargar_hoja_rh.py '/salida/HOJA.xlsx'"
    ... python cargar_hoja_rh.py '/salida/HOJA.xlsx' --aplicar
    ... python cargar_hoja_rh.py --deshacer odoo_cambios_hoja_AAAAMMDD_HHMM.json

  * ENSAYO: no escribe nada. Guarda la huella del plan en
    odoo_plan_hoja.json y en la terminal solo salen cuentas.
  * APLICAR: vuelve a calcular con la hoja y con Odoo de ese momento; si
    algo cambio desde el ensayo, se detiene. Cada cambio se anota con su
    valor de antes en odoo_cambios_hoja_<fecha>.json en el momento de
    hacerlo, asi que un corte a la mitad tambien se puede deshacer.
  * DESHACER: regresa lo anotado, en orden inverso.

La bitacora lleva datos personales: no va a git.
"""
import collections
import hashlib
import json
import os
import re
import sys
import unicodedata
from datetime import datetime

import httpx

AQUI = os.path.dirname(os.path.abspath(__file__))
PLAN = os.path.join(AQUI, "odoo_plan_hoja.json")

LECTURA = {"search_read", "fields_get"}
METODOS = {
    "ensayo": LECTURA,
    "aplicar": LECTURA | {"write"},
    "deshacer": LECTURA | {"write"},
}

PUESTOS = ("personal de seguridad", "security driver")
PLAZAS = ["Ciudad de México", "Guadalajara", "Querétaro", "Monterrey"]
# El Estado de Mexico va como Ciudad de Mexico: la misma zona metropolitana.
ALIAS = {"estado de mexico": "ciudad de mexico", "edomex": "ciudad de mexico",
         "cdmx": "ciudad de mexico"}
DOMINIOS_RAROS = {"gamil.com", "gmial.com", "gmai.com", "gmail.con", "gmail.co",
                  "hotmial.com", "hotmal.com", "hotmail.con", "hotamil.com",
                  "yaho.com", "yahoo.con", "outlok.com", "outlook.con"}
CORREO_VALIDO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# Del titulo de la columna al campo. Por el titulo y no por la posicion:
# si RH inserta una columna, nada se recorre.
COLUMNAS = {
    "no. odoo": "id",
    "nombre": "nombre",
    "plaza": "work_location_id",
    "celular de trabajo": "mobile_phone",
    "correo de trabajo": "work_email",
    "correo personal": "private_email",
    "referencia de empleado": "registration_number",
    "notas de rh": "notas",
}
CAMPOS = ["work_location_id", "mobile_phone", "work_email", "private_email",
          "registration_number"]
ETIQUETA = {"work_location_id": "plaza", "mobile_phone": "celular",
            "work_email": "correo de trabajo", "private_email": "correo personal",
            "registration_number": "referencia"}


class LlaveInvalida(Exception):
    pass


class Odoo:
    def __init__(self, base, llave, modo):
        base = base.strip().rstrip("/")
        self.base = base if base.startswith("http") else "https://" + base
        self.permitidos = METODOS[modo]
        self.http = httpx.Client(timeout=60, headers={
            "Authorization": f"bearer {llave}",
            "Content-Type": "application/json; charset=utf-8",
            "User-Agent": "centauro-hoja-rh/1.0"})
        self.contexto = {"lang": "es_MX"}

    def llamar(self, modelo, metodo, **args):
        if metodo not in self.permitidos:
            raise RuntimeError(f"{metodo} no esta permitido en este modo")
        contexto = {**self.contexto, **args.pop("context", {})}
        r = self.http.post(f"{self.base}/json/2/{modelo}/{metodo}",
                           json={"context": contexto, **args})
        if r.status_code == 200:
            return r.json()
        try:
            error = r.json()
            mensaje = error.get("message") or error.get("name") or r.text
        except ValueError:
            mensaje = r.text
        mensaje = " ".join(str(mensaje).split())[:300]
        if r.status_code == 401:
            raise LlaveInvalida(mensaje)
        raise RuntimeError(f"{modelo}.{metodo} → {r.status_code}: {mensaje}")

    def leer(self, modelo, dominio, campos, **contexto):
        return self.llamar(modelo, "search_read", domain=dominio,
                           fields=campos, order="id", context=contexto) or []


# ------------------------------------------------------------------ ayudas

def normal(texto):
    """Sin acentos, sin mayusculas y sin espacios de sobra."""
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().split())


def texto(valor):
    return "" if valor in (False, None) else " ".join(str(valor).split())


def celda(valor):
    """El texto de una casilla. Excel vuelve numero lo que parece numero:
    un celular llega como 5512345678.0 si alguien le quito el formato."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return texto(valor)


def id_de(valor):
    if not valor:
        return None
    if isinstance(valor, (list, tuple)):
        return valor[0]
    if isinstance(valor, dict):
        return valor.get("id")
    return valor


def nombre_de(valor):
    if isinstance(valor, (list, tuple)) and len(valor) > 1:
        return texto(valor[1])
    if isinstance(valor, dict):
        return texto(valor.get("display_name") or valor.get("name"))
    return ""


def es_de_seguridad(e):
    return any(normal(t).startswith(PUESTOS)
               for t in (nombre_de(e.get("job_id")), e.get("job_title")))


def digitos(t):
    return "".join(c for c in t if c.isdigit())


def igual(campo, nuevo, e):
    """Si lo que dice la hoja es lo que Odoo ya tiene."""
    if campo == "work_location_id":
        return nuevo == id_de(e.get(campo))
    actual = texto(e.get(campo))
    if campo.endswith("email"):
        return normal(nuevo) == normal(actual)
    return nuevo == actual


# ------------------------------------------------------------------ la hoja

def leer_hoja(ruta):
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = libro["Personal"] if "Personal" in libro.sheetnames else None
    if hoja is None:
        for h in libro.worksheets:
            primera = next(h.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if any(normal(c) == "no. odoo" for c in primera):
                hoja = h
                break
    if hoja is None:
        raise RuntimeError("La hoja no trae la pestana «Personal» ni una "
                           "columna «No. Odoo».")
    filas = hoja.iter_rows(values_only=True)
    encabezado = [normal(c) for c in next(filas, ())]
    indice = {}
    for i, titulo in enumerate(encabezado):
        campo = COLUMNAS.get(titulo)
        if campo and campo not in indice:
            indice[campo] = i
    faltan = [t for t, c in COLUMNAS.items() if c not in indice and c != "notas"]
    if faltan:
        raise RuntimeError("A la hoja le faltan columnas: " + ", ".join(faltan))
    salida = []
    for fila in filas:
        if not any(v not in (None, "") for v in fila):
            continue
        salida.append({c: (fila[i] if i < len(fila) else None)
                       for c, i in indice.items()})
    return salida


# ------------------------------------------------------------------ el plan

def calcular(odoo, filas):
    """Lo que habria que escribir, sin escribirlo: (operaciones, resumen)."""
    empleados = odoo.leer("hr.employee", [],
                          ["name", "job_id", "job_title", "company_id"] + CAMPOS)
    seguridad = {e["id"]: e for e in empleados if es_de_seguridad(e)}
    empresas = {id_de(e.get("company_id")) for e in seguridad.values()}
    if len(empresas) != 1:
        raise RuntimeError(f"El personal de seguridad esta en {len(empresas)} "
                           "empresas; se esperaba una. No se hace nada.")
    empresa = empresas.pop()
    odoo.contexto["allowed_company_ids"] = [empresa]

    ubicaciones = odoo.leer("hr.work.location", [], ["name", "company_id"])
    por_nombre = {normal(u["name"]): u["id"] for u in ubicaciones
                  if id_de(u.get("company_id")) in (empresa, None)}
    plazas = {normal(p): por_nombre.get(normal(p)) for p in PLAZAS}
    sin_ubicacion = [p for p in PLAZAS if not plazas[normal(p)]]
    if sin_ubicacion:
        raise RuntimeError("En Odoo no existe la ubicacion de trabajo de: "
                           + ", ".join(sin_ubicacion) + ". No se hace nada.")

    saltadas = collections.Counter()
    no_se_escribe = collections.Counter()
    propuestas, vistos, con_notas = {}, set(), 0

    for fila in filas:
        if celda(fila.get("notas")):
            con_notas += 1
        ident = celda(fila.get("id"))
        if not ident.isdigit():
            saltadas["fila sin No. Odoo"] += 1
            continue
        e = seguridad.get(int(ident))
        if e is None:
            saltadas["no es personal de seguridad en Odoo"] += 1
            continue
        # Primero el nombre y despues si ya vino: una fila movida no le
        # aparta el numero a la fila buena que viene mas abajo.
        if normal(celda(fila.get("nombre"))) != normal(e.get("name")):
            saltadas["el nombre de la fila no es el de Odoo"] += 1
            continue
        if e["id"] in vistos:
            saltadas["la persona viene dos veces"] += 1
            continue
        vistos.add(e["id"])

        # La hoja trae lo que Odoo ya tenia: solo se revisa y se escribe lo
        # que cambia.
        nuevo = {}
        plaza = normal(celda(fila.get("work_location_id")))
        if plaza:
            ubicacion = plazas.get(ALIAS.get(plaza, plaza))
            if ubicacion:
                nuevo["work_location_id"] = ubicacion
            elif plaza != normal(nombre_de(e.get("work_location_id"))):
                no_se_escribe["plaza que no es de las cuatro"] += 1
        celular = celda(fila.get("mobile_phone"))
        if celular and not igual("mobile_phone", celular, e):
            if len(digitos(celular)) >= 10:
                nuevo["mobile_phone"] = celular
            else:
                no_se_escribe["celular de menos de diez digitos"] += 1
        for campo in ("work_email", "private_email"):
            correo = celda(fila.get(campo)).lower()
            if not correo or igual(campo, correo, e):
                continue
            if not CORREO_VALIDO.match(correo):
                no_se_escribe[f"{ETIQUETA[campo]} mal escrito"] += 1
            elif correo.split("@")[-1] in DOMINIOS_RAROS:
                no_se_escribe[f"{ETIQUETA[campo]} con error de dedo"] += 1
            else:
                nuevo[campo] = correo
        referencia = celda(fila.get("registration_number"))
        if referencia:
            nuevo["registration_number"] = referencia
        propuestas[e["id"]] = {c: v for c, v in nuevo.items()
                               if not igual(c, v, e)}

    # El correo con el que entra a la app --el personal, siempre: los de
    # trabajo del personal de seguridad se van a suspender-- no puede ser el
    # de otra persona. Lo que la hoja trae y causaria la repeticion no se
    # escribe.
    def valor(ident, campo):
        if campo in propuestas.get(ident, {}):
            return propuestas[ident][campo]
        return texto(seguridad[ident].get(campo)).lower()

    def acceso_de(ident):
        return valor(ident, "private_email")

    grupos = collections.defaultdict(list)
    for ident in seguridad:
        if acceso_de(ident):
            grupos[acceso_de(ident)].append(ident)
    for correo, ids in grupos.items():
        if len(ids) < 2:
            continue
        for ident in ids:
            p = propuestas.get(ident, {})
            if p.get("private_email") == correo:
                p.pop("private_email")
                no_se_escribe["correo repetido con otra persona"] += 1
    # Lo que sigue repetido despues de esto ya venia asi de Odoo.
    quedan = collections.Counter(acceso_de(i) for i in seguridad if acceso_de(i))
    ya_repetidos = sum(n for n in quedan.values() if n > 1)

    # La referencia es unica en la empresa; se revisa contra todos, con los
    # archivados, porque Odoo tambien los cuenta.
    todas = odoo.leer("hr.employee", [["registration_number", "!=", False]],
                      ["registration_number", "company_id"], active_test=False)
    final = {r["id"]: texto(r["registration_number"]) for r in todas
             if id_de(r.get("company_id")) in (empresa, None)}
    for ident, p in propuestas.items():
        if "registration_number" in p:
            final[ident] = p["registration_number"]
    quienes = collections.defaultdict(set)
    for ident, ref in final.items():
        if ref:
            quienes[normal(ref)].add(ident)
    for ident, p in propuestas.items():
        ref = p.get("registration_number")
        if ref and len(quienes[normal(ref)]) > 1:
            p.pop("registration_number")
            no_se_escribe["referencia repetida con otra persona"] += 1

    ops, por_campo, sin_cambio = [], collections.Counter(), 0
    for ident in sorted(propuestas):
        e, valores = seguridad[ident], propuestas[ident]
        if not valores:
            sin_cambio += 1
            continue
        antes = {c: ((id_de(e.get(c)) if c == "work_location_id" else e.get(c))
                     or False) for c in valores}
        por_campo.update(valores.keys())
        ops.append({"id": ident, "valores": valores, "antes": antes})

    resumen = ["LA HOJA",
               f"  {len(filas)} fila(s) · {len(propuestas)} persona(s) "
               "encontradas en Odoo"]
    if saltadas:
        resumen.append(f"  Saltadas enteras: {sum(saltadas.values())}")
        resumen += [f"    {n:>3}  {m}" for m, n in saltadas.most_common()]
    faltan = len(set(seguridad) - vistos)
    resumen.append(f"  Personal de seguridad que no viene en la hoja: {faltan}")
    if con_notas:
        resumen.append(f"  Filas con notas de RH: {con_notas} (leelas en la "
                       "hoja antes de aplicar)")
    resumen += ["", "LO QUE SE ESCRIBIRIA EN ODOO",
                f"  Personas con algun cambio: {len(ops)}"]
    resumen += [f"    {por_campo[c]:>3}  {ETIQUETA[c]}" for c in CAMPOS
                if por_campo[c]]
    resumen.append(f"  Sin cambio: {sin_cambio}")
    if no_se_escribe:
        resumen += ["", "NO SE ESCRIBE (lo demas de esa persona si)"]
        resumen += [f"    {n:>3}  {m}" for m, n in no_se_escribe.most_common()]
    if ya_repetidos:
        resumen += ["", f"OJO: {ya_repetidos} persona(s) comparten correo de "
                    "acceso en Odoo y la hoja no lo corrige. Centauro las dejara "
                    "pendientes hasta que se arregle."]
    return ops, resumen


def huella(ops):
    return hashlib.sha256(json.dumps(ops, sort_keys=True,
                                     ensure_ascii=False).encode()).hexdigest()


# ------------------------------------------------------------------ aplicar

def aplicar(odoo, filas):
    ops, _ = calcular(odoo, filas)
    if not os.path.exists(PLAN):
        print("No hay ensayo guardado. Corre primero el ensayo (sin --aplicar).")
        return 1
    aprobado = json.load(open(PLAN, encoding="utf-8"))
    if aprobado.get("huella") != huella(ops):
        print("La hoja u Odoo cambiaron desde el ensayo: el plan ya no es el "
              "que se aprobo.\nNo se escribio nada. Corre el ensayo otra vez.")
        return 4
    if not ops:
        print("No hay nada que escribir: Odoo ya tiene lo que dice la hoja.")
        return 0

    ruta = os.path.join(AQUI, f"odoo_cambios_hoja_{datetime.now():%Y%m%d_%H%M}.json")
    bitacora = {"hecho_en": datetime.now().isoformat(timespec="seconds"),
                "odoo": odoo.base, "contexto": odoo.contexto, "cambios": []}

    def anotar(entrada):
        bitacora["cambios"].append(entrada)
        with open(ruta, "w", encoding="utf-8") as f:
            json.dump(bitacora, f, ensure_ascii=False, indent=1)

    hechos = 0
    try:
        for op in ops:
            anotar({"modelo": "hr.employee", "id": op["id"],
                    "antes": op["antes"], "despues": op["valores"]})
            odoo.llamar("hr.employee", "write", ids=[op["id"]],
                        vals=op["valores"])
            hechos += 1
    except (RuntimeError, httpx.HTTPError) as error:
        print(f"\nSe detuvo en la persona {hechos + 1} de {len(ops)}: {error}")
        print(f"Lo que alcanzo a escribir quedo anotado en {os.path.basename(ruta)}.")
        print("Para regresarlo: --deshacer " + os.path.basename(ruta))
        return 5

    # Se comprueba leyendo Odoo otra vez: si el plan sale vacio, quedo.
    restantes, _ = calcular(odoo, filas)
    os.remove(PLAN)
    print(f"Hecho: {hechos} persona(s) actualizadas en Odoo.")
    print("Comprobado en Odoo: " + ("no queda nada por escribir." if not restantes
                                    else f"quedan {len(restantes)} pendientes."))
    print(f"Anotado en {os.path.basename(ruta)}, por si hay que deshacer.")
    return 0 if not restantes else 6


def deshacer(odoo, archivo):
    ruta = archivo if os.path.isabs(archivo) else os.path.join(AQUI, archivo)
    bitacora = json.load(open(ruta, encoding="utf-8"))
    odoo.contexto.update(bitacora.get("contexto", {}))
    regresados = 0
    for c in reversed(bitacora["cambios"]):
        odoo.llamar(c["modelo"], "write", ids=[c["id"]], vals=c["antes"])
        regresados += 1
    print(f"Deshecho: {regresados} persona(s) regresadas a como estaban.")
    return 0


# ------------------------------------------------------------------ main

def main():
    base = os.environ.get("ODOO_BASE", "").strip()
    llave = os.environ.get("ODOO_API_KEY", "").strip()
    if not base or not llave:
        print("Falta ODOO_BASE u ODOO_API_KEY en el .env. Guarda la llave y "
              "vuelve a correr esto.")
        return 1
    args = sys.argv[1:]
    modo = ("aplicar" if "--aplicar" in args
            else "deshacer" if "--deshacer" in args else "ensayo")
    odoo = Odoo(base, llave, modo)
    try:
        if modo == "deshacer":
            i = args.index("--deshacer")
            if i + 1 >= len(args):
                print("Falta el archivo: --deshacer odoo_cambios_hoja_....json")
                return 1
            return deshacer(odoo, args[i + 1])

        hojas = [a for a in args if a.lower().endswith(".xlsx")]
        if len(hojas) != 1:
            print("Falta la hoja: el archivo .xlsx que devolvio RH, entre "
                  "comillas.")
            return 1
        if not os.path.exists(hojas[0]):
            print(f"No encuentro la hoja: {os.path.basename(hojas[0])}. "
                  "Revisa que este en «Claude outputs».")
            return 1
        try:
            filas = leer_hoja(hojas[0])
        except ImportError:
            print("Falta openpyxl: corre el comando con «pip install -q "
                  "openpyxl &&».")
            return 1
        if modo == "aplicar":
            return aplicar(odoo, filas)

        ops, resumen = calcular(odoo, filas)
        print("ENSAYO · no se ha escrito nada en Odoo\n")
        print("\n".join(resumen))
        with open(PLAN, "w", encoding="utf-8") as f:
            json.dump({"huella": huella(ops), "personas": len(ops),
                       "hecho_en": datetime.now().isoformat(timespec="seconds"),
                       "resumen": resumen}, f, ensure_ascii=False, indent=1)
        print("\nSi esta bien, se aplica con --aplicar." if ops
              else "\nNo hay nada que escribir.")
        return 0
    except LlaveInvalida as error:
        print(f"Odoo rechazo la llave: {error}. Crea otra y guardala otra vez.")
        return 2
    except (RuntimeError, httpx.HTTPError) as error:
        print(f"No se pudo: {error}")
        return 3


if __name__ == "__main__":
    sys.exit(main())
